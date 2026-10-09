"""多格式文档解析：TXT/MD、PDF、Word(docx)、Excel(xlsx)、PPT(pptx)、OFD，图片/扫描件走 OCR。

返回 {"text": 全文, "blocks": [{"type":"heading|paragraph|table","text":..}], "tables": [...], "images": n}
"""
import os
import re
import shutil
import subprocess
import zipfile
import xml.etree.ElementTree as ET

import httpx

from .. import config

IMAGE_EXT = {"png", "jpg", "jpeg", "bmp", "tif", "tiff"}
SUPPORTED = {"txt", "md", "pdf", "docx", "xlsx", "pptx", "ofd", "csv"} | IMAGE_EXT


def detect_encoding(raw):
    for enc in ("utf-8", "gbk", "gb18030", "utf-16", "big5"):
        try:
            raw.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def ocr_image(path):
    provider = config.get("ocr.provider", "none")
    if provider == "tesseract":
        cmd = config.get("ocr.tesseract_cmd", "tesseract")
        if shutil.which(cmd):
            r = subprocess.run([cmd, path, "-", "-l", "chi_sim+eng"], capture_output=True, timeout=120)
            return r.stdout.decode("utf-8", "ignore")
    elif provider == "api" and config.get("ocr.api_url"):
        with open(path, "rb") as f:
            r = httpx.post(config.get("ocr.api_url"), files={"file": f}, timeout=120)
        data = r.json()
        return data.get("text") or "\n".join(data.get("lines", []))
    return ""


def _parse_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    blocks, images = [], 0
    for page in reader.pages:
        t = page.extract_text() or ""
        try:
            images += len(page.images)
        except Exception:  # noqa
            pass
        for para in re.split(r"\n\s*\n", t):
            if para.strip():
                blocks.append({"type": "paragraph", "text": para.strip()})
    text = "\n".join(b["text"] for b in blocks)
    if len(text.strip()) < 20 and images:  # 扫描件
        text = "[扫描件，未配置 OCR 时无法提取文字]"
        if config.get("ocr.provider", "none") != "none":
            text = _ocr_pdf_images(reader) or text
        blocks = [{"type": "paragraph", "text": text}]
    return {"text": text, "blocks": blocks, "tables": [], "images": images}


def _ocr_pdf_images(reader):
    import tempfile
    out = []
    for page in reader.pages:
        for img in page.images:
            p = os.path.join(tempfile.gettempdir(), "xx_ocr_" + img.name)
            with open(p, "wb") as f:
                f.write(img.data)
            out.append(ocr_image(p))
            os.remove(p)
    return "\n".join(out)


def _parse_docx(path):
    import docx
    d = docx.Document(path)
    blocks, tables = [], []
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        is_head = p.style is not None and p.style.name and (p.style.name.startswith("Heading") or "标题" in p.style.name)
        blocks.append({"type": "heading" if is_head else "paragraph", "text": p.text.strip()})
    for t in d.tables:
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        tables.append(rows)
        blocks.append({"type": "table", "text": "\n".join(" | ".join(r) for r in rows)})
    images = len(d.inline_shapes)
    return {"text": "\n".join(b["text"] for b in blocks), "blocks": blocks, "tables": tables, "images": images}


def _parse_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    blocks, tables = [], []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            if any(v is not None for v in r):
                rows.append(["" if v is None else str(v) for v in r])
        tables.append(rows)
        blocks.append({"type": "heading", "text": "工作表：" + ws.title})
        blocks.append({"type": "table", "text": "\n".join(" | ".join(r) for r in rows)})
    return {"text": "\n".join(b["text"] for b in blocks), "blocks": blocks, "tables": tables, "images": 0}


def _parse_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    blocks, images = [], 0
    for i, slide in enumerate(prs.slides, 1):
        blocks.append({"type": "heading", "text": "第%d页" % i})
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                blocks.append({"type": "paragraph", "text": shape.text_frame.text.strip()})
            if shape.shape_type == 13:
                images += 1
    return {"text": "\n".join(b["text"] for b in blocks), "blocks": blocks, "tables": [], "images": images}


def _parse_ofd(path):
    """OFD 为 zip 包，正文位于 Doc_x/Pages/Page_x/Content.xml 的 TextCode 节点"""
    blocks = []
    with zipfile.ZipFile(path) as z:
        pages = sorted([n for n in z.namelist() if re.search(r"Pages/Page_\d+/Content\.xml$", n)],
                       key=lambda n: int(re.search(r"Page_(\d+)", n).group(1)))
        for name in pages:
            root = ET.fromstring(z.read(name))
            line, last_y = [], None
            for el in root.iter():
                if el.tag.endswith("TextObject"):
                    boundary = el.get("Boundary", "0 0 0 0").split()
                    y = float(boundary[1]) if len(boundary) > 1 else 0
                    if last_y is not None and abs(y - last_y) > 2 and line:
                        blocks.append({"type": "paragraph", "text": "".join(line)})
                        line = []
                    last_y = y
                    for tc in el:
                        if tc.tag.endswith("TextCode") and tc.text:
                            line.append(tc.text)
            if line:
                blocks.append({"type": "paragraph", "text": "".join(line)})
    return {"text": "\n".join(b["text"] for b in blocks), "blocks": blocks, "tables": [], "images": 0}


def _parse_text(path):
    with open(path, "rb") as f:
        raw = f.read()
    text = raw.decode(detect_encoding(raw), "ignore")
    blocks = [{"type": "heading" if p.startswith("#") else "paragraph", "text": p.strip()}
              for p in re.split(r"\n\s*\n", text) if p.strip()]
    return {"text": text, "blocks": blocks, "tables": [], "images": 0}


def parse_file(path):
    ext = path.rsplit(".", 1)[-1].lower()
    if ext in ("txt", "md", "csv"):
        return _parse_text(path)
    if ext == "pdf":
        return _parse_pdf(path)
    if ext == "docx":
        return _parse_docx(path)
    if ext == "xlsx":
        return _parse_xlsx(path)
    if ext == "pptx":
        return _parse_pptx(path)
    if ext == "ofd":
        return _parse_ofd(path)
    if ext in IMAGE_EXT:
        t = ocr_image(path) or "[图片，未配置 OCR]"
        return {"text": t, "blocks": [{"type": "paragraph", "text": t}], "tables": [], "images": 1}
    raise ValueError("不支持的文件格式: " + ext)


# ---------------- 文本预处理 ----------------
_GARBLED = re.compile(r"[�\u0000-\u0008\u000b\u000c\u000e-\u001f]")


def clean_text(text):
    text = _GARBLED.sub("", text)
    text = text.replace("　", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def is_low_quality(chunk):
    """过滤无效信息：过短、乱码比例高、纯符号/数字"""
    s = chunk.strip()
    if len(s) < 10:
        return True
    valid = len(re.findall(r"[一-龥A-Za-z]", s))
    return valid / max(len(s), 1) < 0.3


def split_chunks(text, size=500, overlap=80):
    """按段落结构切分，段落过长时滑动切分"""
    paras = [p.strip() for p in re.split(r"\n+", text) if p.strip()]
    chunks, buf = [], ""
    for p in paras:
        if p.startswith("#") and buf:  # Markdown 标题处切分，保持段落结构
            chunks.append(buf)
            buf = ""
        if len(buf) + len(p) + 1 <= size:
            buf = (buf + "\n" + p) if buf else p
            continue
        if buf:
            chunks.append(buf)
        while len(p) > size:
            chunks.append(p[:size])
            p = p[size - overlap:]
        buf = p
    if buf:
        chunks.append(buf)
    return chunks
